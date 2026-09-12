"""Generate a Word (.docx) version of Bhupraj Tamang's NV1 Security Officer CV."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
DARK_GREY = RGBColor(0x33, 0x33, 0x33)


def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def add_section_heading(doc, text):
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(4)
    run = heading.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = NAVY
    pPr = heading._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return heading


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.size = Pt(10.5)
        run.font.color.rgb = DARK_GREY
    return p


doc = Document()

# Base document styling
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.font.color.rgb = DARK_GREY

for section in doc.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

# Name header
name_p = doc.add_paragraph()
name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
name_p.paragraph_format.space_after = Pt(2)
name_run = name_p.add_run("BHUPRAJ TAMANG")
name_run.bold = True
name_run.font.size = Pt(24)
name_run.font.color.rgb = NAVY

contact_p = doc.add_paragraph()
contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_p.paragraph_format.space_after = Pt(2)
contact_run = contact_p.add_run("Darwin, NT 0800  |  0447 481 904  |  [Your Email Address]")
contact_run.font.size = Pt(10.5)
contact_run.font.color.rgb = DARK_GREY

avail_p = doc.add_paragraph()
avail_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
avail_p.paragraph_format.space_after = Pt(6)
avail_run = avail_p.add_run(
    "Full Australian Working Rights  |  Available for All Shifts "
    "(Day / Night / Weekend / 4-on 4-off Rotation)"
)
avail_run.italic = True
avail_run.font.size = Pt(9.5)
avail_run.font.color.rgb = DARK_GREY

# Divider line under header
divider_p = doc.add_paragraph()
divider_p.paragraph_format.space_after = Pt(2)
pPr = divider_p._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "18")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), "1F3A5F")
pBdr.append(bottom)
pPr.append(pBdr)

# Professional Summary
add_section_heading(doc, "Professional Summary")
summary_p = doc.add_paragraph()
summary_p.paragraph_format.space_after = Pt(6)
summary_run = summary_p.add_run(
    "Reliable and safety-focused Security Officer with hands-on customer service experience "
    "and a full suite of current security and first aid qualifications, including a "
    "Certificate II in Security Operations (CPP20218), Security Officer Licence, and Crowd "
    "Controller Licence issued under the Northern Territory Private Security Act. Comfortable "
    "working independently, following strict procedures, and maintaining a calm, professional "
    "presence in fast-paced, customer-facing environments. Physically fit, punctual, and "
    "available for a full range of shifts including day, night, weekend, and 4-on/4-off "
    "rotations. Seeking to bring strong attention to detail, risk-awareness, and service "
    "excellence to the NV1 Security Officer role with Millennium Services Group in Darwin."
)
summary_run.font.size = Pt(10.5)
summary_run.font.color.rgb = DARK_GREY

note_p = doc.add_paragraph()
note_p.paragraph_format.space_after = Pt(4)
note_run = note_p.add_run(
    "Note: This role requires an active NV1 security clearance. "
    "[Insert your NV1 clearance status here.]"
)
note_run.italic = True
note_run.font.size = Pt(9.5)
note_run.font.color.rgb = RGBColor(0x80, 0x00, 0x00)

# Licences & Certifications
add_section_heading(doc, "Licences & Certifications")

cert_data = [
    ("Security Officer Licence", "Northern Territory Private Security Act", "DLSC6551 — Expires 03/06/2028"),
    ("Crowd Controller Licence", "Northern Territory Private Security Act", "DLSC6551 — Expires 03/06/2028"),
    ("Certificate II in Security Operations (CPP20218)", "Integrated Training", "—"),
    ("CPCWHS1001 – Prepare to Work Safely in the Construction Industry", "Integrated Training", "—"),
    ("HLTAID011 – Provide First Aid", "Integrated Training", "—"),
    ("HLTAID009 – Provide Cardiopulmonary Resuscitation (CPR)", "Integrated Training", "—"),
    ("HLTAID010 – Provide Basic Emergency Life Support", "Integrated Training", "—"),
    ("Responsible Service of Alcohol (RSA)", "Integrated Training", "—"),
]

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = True

hdr_cells = table.rows[0].cells
headers = ["Licence / Certificate", "Issuing Body", "Licence No. / Expiry"]
for i, h in enumerate(headers):
    hdr_cells[i].text = ""
    p = hdr_cells[i].paragraphs[0]
    r = p.add_run(h)
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_cell_shading(hdr_cells[i], "1F3A5F")

for cert, body, expiry in cert_data:
    row_cells = table.add_row().cells
    for i, val in enumerate([cert, body, expiry]):
        row_cells[i].text = ""
        p = row_cells[i].paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(9.5)
        r.font.color.rgb = DARK_GREY

widths = [Inches(3.5), Inches(2.1), Inches(1.4)]
for row in table.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = w

# Core Skills
add_section_heading(doc, "Core Skills")
skills = [
    "Access control & visitor/contractor log management",
    "CCTV monitoring and footage review",
    "Walk-through and handheld metal detector operation",
    "Incident, emergency, and alarm response",
    "Risk identification and assessment",
    "Incident report writing and accurate record-keeping",
    "Conflict de-escalation and customer service",
    "Site patrols and physical presence duties",
    "Working independently and adhering to procedures",
    "Strong attention to detail and reliability",
]
for skill in skills:
    add_bullet(doc, skill)

# Work Experience
add_section_heading(doc, "Work Experience")

role_p = doc.add_paragraph()
role_p.paragraph_format.space_after = Pt(0)
role_run = role_p.add_run("Team Member — Coles Supermarkets")
role_run.bold = True
role_run.font.size = Pt(11)
role_run.font.color.rgb = NAVY

date_p = doc.add_paragraph()
date_p.paragraph_format.space_after = Pt(4)
date_run = date_p.add_run("Darwin, NT  |  January 2026 – Present")
date_run.italic = True
date_run.font.size = Pt(9.5)
date_run.font.color.rgb = DARK_GREY

experience_bullets = [
    "Delivered consistent, friendly customer service in a high-traffic retail environment, "
    "handling customer queries, complaints, and requests professionally.",
    "Monitored the shop floor for safety hazards, suspicious behaviour, and stock-loss risks, "
    "escalating concerns to management or security in line with store procedures.",
    "Operated point-of-sale systems and handled cash and eftpos transactions accurately, "
    "following strict company protocols.",
    "Maintained a calm, approachable, and professional presence, including during busy periods "
    "and when managing difficult customer interactions.",
    "Worked flexibly across a variety of shifts, including early mornings, evenings, and "
    "weekends, consistently maintaining a reliable and punctual attendance record.",
    "Followed workplace health and safety procedures and completed all required compliance "
    "training.",
]
for bullet in experience_bullets:
    add_bullet(doc, bullet)

# Availability
add_section_heading(doc, "Availability")
add_bullet(doc, "Available for any shift: day, night, weekend, and 4-on/4-off rotations")
add_bullet(doc, "Full Australian working rights")
add_bullet(doc, "Able to start immediately")

# Referees
add_section_heading(doc, "Referees")
ref_p = doc.add_paragraph()
ref_run = ref_p.add_run("Available upon request.")
ref_run.font.size = Pt(10.5)
ref_run.font.color.rgb = DARK_GREY

doc.save("/workspace/CV_Bhupraj_Tamang_NV1_Security_Officer.docx")
print("Saved DOCX successfully.")
